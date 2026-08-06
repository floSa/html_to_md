"""Fixtures partagées : profils d'extraction et documents bureautiques.

Les documents Word et PowerPoint sont fabriqués à la volée plutôt que
versionnés : un binaire dans le dépôt se relit mal et se modifie encore
moins bien quand un test doit évoluer.
"""

from __future__ import annotations

import random
import struct
import zlib
from pathlib import Path

import pytest

from html_to_md.extract import Profile, load_profiles

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def profiles() -> list[Profile]:
    """Profils réels du dépôt, pour tester ce qui tourne en production."""
    return load_profiles(ROOT / "config" / "selectors.yaml")


@pytest.fixture
def no_profiles() -> list[Profile]:
    """Aucun profil : force le mode générique (conteneur sémantique / repli)."""
    return []


def make_png(width: int, height: int, seed: int = 0) -> bytes:
    """PNG de bruit aléatoire : incompressible, donc d'une taille réaliste.

    Une image unie compresserait à quelques centaines d'octets et passerait
    pour une icône, ce qui fausserait les tests sur le seuil de taille.
    """
    rng = random.Random(seed)
    raw = b"".join(
        b"\x00" + bytes(rng.randrange(256) for _ in range(width * 3))
        for _ in range(height)
    )

    def chunk(tag: bytes, payload: bytes) -> bytes:
        body = tag + payload
        return struct.pack(">I", len(payload)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


@pytest.fixture
def docx_factory(tmp_path: Path):
    """Fabrique un .docx contenant un titre, deux images et un tableau."""

    def build(name: str = "Rapport de test") -> Path:
        from docx import Document
        from docx.shared import Inches

        big = tmp_path / "schema.png"
        small = tmp_path / "vignette.png"
        big.write_bytes(make_png(120, 120, seed=1))
        small.write_bytes(make_png(20, 20, seed=2))

        document = Document()
        document.add_heading("Titre du rapport", level=1)
        document.add_paragraph("Un paragraphe avant image.")
        document.add_picture(str(big), width=Inches(2))
        document.add_paragraph("Une petite vignette de contenu :")
        document.add_picture(str(small), width=Inches(0.3))

        table = document.add_table(rows=3, cols=3)
        rows = [
            ["Produit", "Prix", "Stock"],
            ["Alpha", "12,50", "3"],
            ["Beta", "7,00", "15"],
        ]
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
        document.add_paragraph("Texte après tableau.")

        path = tmp_path / f"{name}.docx"
        document.save(str(path))
        return path

    return build


@pytest.fixture
def pptx_factory(tmp_path: Path):
    """Fabrique un .pptx d'une diapositive titrée."""

    def build(name: str = "Ma presentation") -> Path:
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Titre de la présentation"
        slide.placeholders[1].text = "Premier point"

        path = tmp_path / f"{name}.pptx"
        presentation.save(str(path))
        return path

    return build
